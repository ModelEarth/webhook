from flask import Flask, request, jsonify
from docx import Document
import os
import re

application = Flask(__name__)

placeholders={
    '#startDate#': '26b40600',
    '#schoolDegree#': '11c2cd54', #optional
    '#degreeDate#': '0b03fd30', #optional
    '#name#': '6a2f8ab5',
    '#optDepartment#': '01bcc04b', #optional
    '#optContact#': '3ab50c1d', #optional
    '#workingHours#': '0c8671f9',
    '#github#': '51a6d3df',
    '#email#': '61c73e81',
    '#endDate#': '31e5f166', #optional
    '#phone#': '086a67d6' #optional
}

@application.route('/signup', methods=['POST'])
def webhook():
    data = request.json
    doc = Document('offer_letter_template.docx')

    # extract answers
    answers = {}
    for question_id, answer_data in data['answers'].items():
        answers[question_id] = [answer['value'] for answer in answer_data['textAnswers']['answers']]

    # replace placeholders
    pattern = re.compile(r'#\w+#')
    for paragraph in doc.paragraphs:
        for match in re.findall(pattern, paragraph.text):
            if match in placeholders:
                # check optional fields in the form
                to_replace = ''
                if placeholders[match] in answers:
                    to_replace = answers[placeholders[match]][0]
                paragraph.text = paragraph.text.replace(match, to_replace)
  
    doc_path = os.path.join('generated_docs', f"{answers['6a2f8ab5'][0]}_Offer_Letter.docx")
    if not os.path.exists('generated_docs'):
        os.makedirs('generated_docs')
    doc.save(doc_path)

    return jsonify({'status': 'success', 'documentPath': doc_path}), 200

if __name__ == '__main__':
    application.run(debug=True, host='127.0.0.1', port=8000)